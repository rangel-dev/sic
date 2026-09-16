"""Suíte de testes do SegmentadoEngine.

Não existia nenhum teste dedicado a este engine antes do BRD-011. Cobre o
parsing/scan (comportamento pré-existente, sem nenhuma mudança de
comportamento) e, em especial, prova que o fallback do BRD-011 (título de
coluna sem espaço/underscore, ex.: 'PORSEGMENTADO' no lugar de
'POR SEGMENTADO') é estritamente aditivo: toda entrada que já funcionava com
o título de hoje continua produzindo exatamente o mesmo resultado, e o
fallback só entra em ação quando a busca exata não encontra nada.
"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import openpyxl
from lxml import etree

from src.core.segmentado_engine import PRICEBOOK_NS, SegmentadoEngine as E
from src.core.segmentado_engine import SegmentedSku


def _make_workbook(tmp_path: Path, sheets: list[tuple[str, list[list], bool]], filename: str = "grade.xlsx") -> Path:
    """sheets: lista de (nome_aba, linhas, oculta) — mesmo padrão usado em
    tests/test_sync_engine_parse_excel.py."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove a aba padrão "Sheet"
    for name, rows, hidden in sheets:
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        if hidden:
            ws.sheet_state = "hidden"
    if all(hidden for _, _, hidden in sheets):
        wb.create_sheet("_dummy_visible")
    path = tmp_path / filename
    wb.save(path)
    return path


# ─── _normalize (comportamento pré-existente) ──────────────────────────────

class TestNormalize:
    def test_upper_e_strip(self):
        assert E._normalize("  por segmentado  ") == "POR SEGMENTADO"

    def test_colapsa_espacos_multiplos(self):
        assert E._normalize("POR   SEGMENTADO") == "POR SEGMENTADO"

    def test_none_vira_string_vazia(self):
        assert E._normalize(None) == ""


# ─── _find_header_row — comportamento de hoje (título com espaço) ─────────

class TestFindHeaderRowTituloComEspacoDeHoje:
    """Este grupo prova que o BRD-011 não mudou nada do que já funcionava."""

    def test_por_segmentado_com_espaco_simples(self):
        rows = [("SKU", "POR SEGMENTADO", "OUTRA")]
        found = E._find_header_row(rows)
        assert found == (0, {"SKU": 0, "POR SEGMENTADO": 1, "OUTRA": 2}, "POR SEGMENTADO")

    def test_por_origem_com_espaco_simples(self):
        rows = [("SKU", "POR ORIGEM")]
        found = E._find_header_row(rows)
        assert found is not None
        assert found[2] == "POR ORIGEM"

    def test_espacos_multiplos_continuam_colapsando(self):
        rows = [("SKU", "POR   SEGMENTADO")]
        found = E._find_header_row(rows)
        assert found is not None
        assert found[2] == "POR SEGMENTADO"

    def test_nenhuma_coluna_alvo_nao_casa(self):
        rows = [("SKU", "QUALQUER OUTRA COISA")]
        assert E._find_header_row(rows) is None

    def test_header_fora_da_janela_de_scan_nao_e_encontrado(self):
        # _HEADER_SCAN_ROWS = 15 — cabeçalho na 16ª linha não deve ser achado.
        rows = [("nada",)] * 15 + [("SKU", "POR SEGMENTADO")]
        assert E._find_header_row(rows) is None


# ─── _find_header_row — fallback do BRD-011 (título futuro sem espaço) ────

class TestFindHeaderRowFallbackBRD011:
    """Título futuro (espaço/underscore removido) só é aceito quando a busca
    exata (título com espaço, o de hoje) não encontrar nada."""

    def test_sem_espaco_algum(self):
        rows = [("SKU", "PORSEGMENTADO")]
        found = E._find_header_row(rows)
        assert found is not None
        assert found[2] == "PORSEGMENTADO"  # chave real do header_map, não a constante

    def test_com_underscore(self):
        rows = [("SKU", "POR_SEGMENTADO")]
        found = E._find_header_row(rows)
        assert found is not None
        assert found[2] == "POR_SEGMENTADO"

    def test_por_origem_sem_espaco(self):
        rows = [("SKU", "PORORIGEM")]
        found = E._find_header_row(rows)
        assert found is not None
        assert found[2] == "PORORIGEM"

    def test_matched_header_e_chave_valida_do_header_map(self):
        """O valor retornado é usado depois em header_map.get(matched_header)
        para achar a coluna de preço — precisa ser uma chave real."""
        rows = [("SKU", "POR_SEGMENTADO")]
        found = E._find_header_row(rows)
        assert found is not None
        _, header_map, matched = found
        assert header_map.get(matched) == 1

    def test_titulo_com_espaco_tem_prioridade_sobre_o_fallback(self):
        """Se as duas grafias existissem juntas, a exata vence — o fallback
        só roda quando a busca exata falha."""
        rows = [("POR SEGMENTADO", "PORSEGMENTADO")]
        found = E._find_header_row(rows)
        assert found is not None
        _, _, matched = found
        assert matched == "POR SEGMENTADO"

    def test_titulo_nao_relacionado_nao_casa_por_acidente(self):
        rows = [("SKU", "DESCRICAO DO PRODUTO")]
        assert E._find_header_row(rows) is None


# ─── _extract_total_informado ──────────────────────────────────────────────

class TestExtractTotalInformado:
    def test_total_skus_com_espaco_de_hoje(self):
        rows = [(None,), (None,), ("TOTAL SKUS", 42)]
        assert E._extract_total_informado(rows) == 42

    def test_total_skus_sem_espaco_fallback_brd011(self):
        rows = [(None,), (None,), ("TOTALSKUS", 42)]
        assert E._extract_total_informado(rows) == 42

    def test_valor_nao_numerico_e_ignorado(self):
        rows = [(None,), (None,), ("TOTAL SKUS", "quarenta e dois")]
        assert E._extract_total_informado(rows) is None

    def test_bool_nao_conta_como_numero(self):
        rows = [(None,), (None,), ("TOTAL SKUS", True)]
        assert E._extract_total_informado(rows) is None

    def test_linhas_insuficientes_retorna_none(self):
        assert E._extract_total_informado([(None,)]) is None

    def test_rotulo_nao_relacionado_retorna_none(self):
        rows = [(None,), (None,), ("OUTRO ROTULO", 42)]
        assert E._extract_total_informado(rows) is None


# ─── _parse_price_cell (comportamento pré-existente) ───────────────────────

class TestParsePriceCell:
    def test_int(self):
        assert E._parse_price_cell(30) == 30.0

    def test_float(self):
        assert E._parse_price_cell(29.9) == 29.9

    def test_string_com_virgula_decimal(self):
        assert E._parse_price_cell("29,90") == 29.9

    def test_string_com_milhar_e_virgula(self):
        assert E._parse_price_cell("1.234,56") == 1234.56

    def test_string_com_prefixo_moeda(self):
        assert E._parse_price_cell("R$ 30,00") == 30.0

    def test_string_com_nbsp(self):
        assert E._parse_price_cell("R$\xa030,00") == 30.0

    def test_string_vazia_retorna_none(self):
        assert E._parse_price_cell("") is None

    def test_none_retorna_none(self):
        assert E._parse_price_cell(None) is None

    def test_bool_retorna_none(self):
        assert E._parse_price_cell(True) is None

    def test_string_invalida_retorna_none(self):
        assert E._parse_price_cell("abc") is None


# ─── _pick_product_id_column (comportamento pré-existente) ────────────────

class TestPickProductIdColumn:
    def test_escolhe_coluna_com_mais_skus_validos(self):
        rows = [
            ("header1", "header2"),
            ("NATBRA-001", "texto qualquer"),
            ("NATBRA-002", "texto qualquer"),
            ("NATBRA-003", "NATBRA-999"),
        ]
        assert E._pick_product_id_column(rows, 0) == 0

    def test_sem_nenhum_sku_retorna_none(self):
        rows = [("header",), ("abc", "def")]
        assert E._pick_product_id_column(rows, 0) is None


# ─── _extract_periodo (comportamento pré-existente) ────────────────────────

class TestExtractPeriodo:
    def test_periodo_valido(self):
        rows = [(None, date(2026, 1, 1), date(2026, 1, 31))]
        assert E._extract_periodo(rows) == (date(2026, 1, 1), date(2026, 1, 31))

    def test_datetime_e_convertido_para_date(self):
        rows = [(None, datetime(2026, 1, 1, 10, 0), datetime(2026, 1, 31, 10, 0))]
        assert E._extract_periodo(rows) == (date(2026, 1, 1), date(2026, 1, 31))

    def test_fim_antes_do_inicio_retorna_none(self):
        rows = [(None, date(2026, 1, 31), date(2026, 1, 1))]
        assert E._extract_periodo(rows) is None

    def test_sem_datas_retorna_none(self):
        rows = [(None, "texto", "texto")]
        assert E._extract_periodo(rows) is None

    def test_linha_insuficiente_retorna_none(self):
        assert E._extract_periodo([]) is None


# ─── _extract_lp_label (comportamento pré-existente) ───────────────────────

class TestExtractLpLabel:
    def test_lp_encontrado(self):
        rows = [(None,), ("LP", "Campanha Natal")]
        assert E._extract_lp_label(rows) == "Campanha Natal"

    def test_sem_lp_retorna_none(self):
        rows = [(None,), ("OUTRO ROTULO", "valor")]
        assert E._extract_lp_label(rows) is None

    def test_valor_vazio_retorna_none(self):
        rows = [(None,), ("LP", None)]
        assert E._extract_lp_label(rows) is None

    def test_linha_insuficiente_retorna_none(self):
        assert E._extract_lp_label([(None,)]) is None


# ─── _extract_rows (comportamento pré-existente) ───────────────────────────

class TestExtractRows:
    def test_extrai_sku_com_preco_valido(self):
        rows = [
            ("SKU", "POR SEGMENTADO"),
            ("NATBRA-001", "30,00"),
        ]
        skus, stats = E._extract_rows(rows, 0, 0, 1, "NATBRA-")
        assert skus == [SegmentedSku(product_id="NATBRA-001", price=30.0)]
        assert stats["scanned"] == 1
        assert stats["nat"] == 1

    def test_sku_de_outra_marca_e_filtrado(self):
        rows = [
            ("SKU", "POR SEGMENTADO"),
            ("AVNBRA-001", "30,00"),
        ]
        skus, stats = E._extract_rows(rows, 0, 0, 1, "NATBRA-")
        assert skus == []
        assert stats["foreign"] == 1

    def test_sem_preco_e_ignorado(self):
        rows = [
            ("SKU", "POR SEGMENTADO"),
            ("NATBRA-001", None),
        ]
        skus, stats = E._extract_rows(rows, 0, 0, 1, "NATBRA-")
        assert skus == []
        assert stats["no_price"] == 1

    def test_preco_zero_e_ignorado(self):
        rows = [
            ("SKU", "POR SEGMENTADO"),
            ("NATBRA-001", 0),
        ]
        skus, stats = E._extract_rows(rows, 0, 0, 1, "NATBRA-")
        assert skus == []
        assert stats["no_price"] == 1

    def test_sku_duplicado_ultima_ocorrencia_vence(self):
        rows = [
            ("SKU", "POR SEGMENTADO"),
            ("NATBRA-001", "10,00"),
            ("NATBRA-001", "20,00"),
        ]
        skus, stats = E._extract_rows(rows, 0, 0, 1, "NATBRA-")
        assert skus == [SegmentedSku(product_id="NATBRA-001", price=20.0)]
        assert stats["dup"] == 1

    def test_para_apos_50_linhas_vazias_consecutivas(self):
        rows = [("SKU", "POR SEGMENTADO")] + [(None, None)] * 60 + [("NATBRA-999", "10,00")]
        skus, _ = E._extract_rows(rows, 0, 0, 1, "NATBRA-")
        assert skus == []  # corte acontece antes de chegar no SKU final


# ─── scan_workbook — fluxo completo, comportamento de hoje ────────────────

class TestScanWorkbookTituloComEspacoDeHoje:
    def test_aba_com_por_segmentado_e_reconhecida(self, tmp_path):
        path = _make_workbook(tmp_path, [
            ("LISTA_01", [
                ["SKU", "POR SEGMENTADO"],
                ["NATBRA-001", "30,00"],
                ["NATBRA-002", "40,00"],
            ], False),
        ])
        result = E.scan_workbook(str(path), expected_brand="natura")
        assert result.error is None
        assert len(result.candidates) == 1
        cand = result.candidates[0]
        assert cand.matched_header == "POR SEGMENTADO"
        assert {s.product_id for s in cand.rows} == {"NATBRA-001", "NATBRA-002"}

    def test_aba_oculta_e_ignorada(self, tmp_path):
        path = _make_workbook(tmp_path, [
            ("LISTA_01", [
                ["SKU", "POR SEGMENTADO"],
                ["NATBRA-001", "30,00"],
            ], True),
        ])
        result = E.scan_workbook(str(path), expected_brand="natura")
        assert result.candidates == []
        assert result.error is not None

    def test_nenhuma_coluna_alvo_gera_mensagem_de_erro(self, tmp_path):
        path = _make_workbook(tmp_path, [
            ("LISTA_01", [["SKU", "OUTRA COLUNA"]], False),
        ])
        result = E.scan_workbook(str(path), expected_brand="natura")
        assert result.candidates == []
        assert result.error is not None
        assert "Nenhuma aba" in result.error


# ─── scan_workbook — fallback do BRD-011 é aditivo ─────────────────────────

class TestScanWorkbookFallbackBRD011:
    """Prova, de ponta a ponta, que reconhecer o título futuro sem espaço
    não altera em nada o resultado do título de hoje: as duas planilhas
    (só muda a grafia do cabeçalho) produzem exatamente o mesmo resultado."""

    def test_titulo_sem_espaco_produz_resultado_identico_ao_titulo_de_hoje(self, tmp_path):
        rows_hoje = [
            ["SKU", "POR SEGMENTADO"],
            ["NATBRA-001", "30,00"],
            ["NATBRA-002", "40,00"],
        ]
        rows_futuro = [
            ["SKU", "PORSEGMENTADO"],
            ["NATBRA-001", "30,00"],
            ["NATBRA-002", "40,00"],
        ]
        path_hoje = _make_workbook(tmp_path, [("LISTA_01", rows_hoje, False)], filename="grade_hoje.xlsx")
        path_futuro = _make_workbook(tmp_path, [("LISTA_01", rows_futuro, False)], filename="grade_futuro.xlsx")

        result_hoje = E.scan_workbook(str(path_hoje), expected_brand="natura")
        result_futuro = E.scan_workbook(str(path_futuro), expected_brand="natura")

        assert result_hoje.error is None
        assert result_futuro.error is None
        assert len(result_hoje.candidates) == len(result_futuro.candidates) == 1

        skus_hoje = sorted(s.product_id for s in result_hoje.candidates[0].rows)
        skus_futuro = sorted(s.product_id for s in result_futuro.candidates[0].rows)
        assert skus_hoje == skus_futuro == ["NATBRA-001", "NATBRA-002"]

    def test_titulo_com_underscore_tambem_e_reconhecido(self, tmp_path):
        path = _make_workbook(tmp_path, [
            ("LISTA_01", [
                ["SKU", "POR_ORIGEM"],
                ["NATBRA-001", "30,00"],
            ], False),
        ])
        result = E.scan_workbook(str(path), expected_brand="natura")
        assert result.error is None
        assert len(result.candidates) == 1
        assert result.candidates[0].matched_header == "POR_ORIGEM"


# ─── compute_online_window (comportamento pré-existente) ──────────────────

class TestComputeOnlineWindow:
    def test_converte_brt_para_utc(self):
        start = datetime(2026, 1, 1, 9, 0)
        end = datetime(2026, 1, 2, 18, 30)
        online_from, online_to = E.compute_online_window(start, end)
        assert online_from == "2026-01-01T12:00:00.000Z"
        assert online_to == "2026-01-02T21:30:00.000Z"

    def test_zera_segundos_e_microssegundos(self):
        start = datetime(2026, 1, 1, 9, 0, 45, 123456)
        end = datetime(2026, 1, 1, 10, 0, 0)
        online_from, _ = E.compute_online_window(start, end)
        assert online_from == "2026-01-01T12:00:00.000Z"


# ─── build_xml (comportamento pré-existente) ───────────────────────────────

class TestBuildXml:
    def test_estrutura_basica(self):
        xml_bytes = E.build_xml(
            pricebook_id="pb-123",
            parent_id="pb-parent",
            online_from="2026-01-01T00:00:00.000Z",
            online_to="2026-01-02T00:00:00.000Z",
            skus=[SegmentedSku(product_id="NATBRA-001", price=30.0)],
        )
        root = etree.fromstring(xml_bytes)
        ns = f"{{{PRICEBOOK_NS}}}"
        pb = root.find(f"{ns}pricebook")
        assert pb is not None

        header = pb.find(f"{ns}header")
        assert header.get("pricebook-id") == "pb-123"
        assert header.find(f"{ns}parent").text == "pb-parent"
        assert header.find(f"{ns}online-flag").text == "true"

        price_table = pb.find(f"{ns}price-tables/{ns}price-table")
        assert price_table.get("product-id") == "NATBRA-001"
        assert price_table.find(f"{ns}amount").text == "30.00"

    def test_display_name_opcional(self):
        xml_com_nome = E.build_xml(
            pricebook_id="pb-1", parent_id="pb-parent",
            online_from="2026-01-01T00:00:00.000Z", online_to="2026-01-02T00:00:00.000Z",
            skus=[], display_name="Campanha X",
        )
        xml_sem_nome = E.build_xml(
            pricebook_id="pb-1", parent_id="pb-parent",
            online_from="2026-01-01T00:00:00.000Z", online_to="2026-01-02T00:00:00.000Z",
            skus=[],
        )
        ns = f"{{{PRICEBOOK_NS}}}"
        assert etree.fromstring(xml_com_nome).find(f"{ns}pricebook/{ns}header/{ns}display-name") is not None
        assert etree.fromstring(xml_sem_nome).find(f"{ns}pricebook/{ns}header/{ns}display-name") is None
