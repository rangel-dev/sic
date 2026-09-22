"""Suíte de testes do segmentada_registry (BRD-012, P1).

Cobre as funções puras (montagem/leitura de ID, cálculo de status, regra de
reconhecimento da Etapa 3) e o armazenamento local (`JsonFileRegistryStore`,
`InMemoryRegistryStore`). Nenhum destes caminhos depende de tela ou de rede.
"""
from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest

import openpyxl

from src.core.segmentada_registry import (
    InMemoryRegistryStore,
    JsonFileRegistryStore,
    MatchConfidence,
    SegmentadaRecord,
    SyncedFolderRegistryStore,
    build_pricebook_id,
    compute_status,
    merge_by_pricebook_id,
    normalize_sheet_name,
    parse_runrun_id,
    resolve,
    write_xlsx_snapshot,
)


def _record(**overrides) -> SegmentadaRecord:
    base = dict(
        pricebook_id="NAT-RR1111",
        sheet_name="LISTA_48",
        brand="natura",
        loja="natura",
        lp_label="Favoritos RR18/21",
    )
    base.update(overrides)
    return SegmentadaRecord(**base)


# ─── normalize_sheet_name ───────────────────────────────────────────────

class TestNormalizeSheetName:
    def test_upper_e_strip(self):
        assert normalize_sheet_name("  lista_48  ") == "LISTA_48"

    def test_colapsa_espacos_multiplos(self):
        assert normalize_sheet_name("LISTA   48") == "LISTA 48"

    def test_none_vira_string_vazia(self):
        assert normalize_sheet_name(None) == ""


# ─── build_pricebook_id / parse_runrun_id ───────────────────────────────

class TestBuildPricebookId:
    def test_natura(self):
        assert build_pricebook_id("natura", "1111") == "NAT-RR1111"

    def test_avon(self):
        assert build_pricebook_id("avon", "2222") == "AVN-RR2222"

    def test_ml_nao_gera_id_automatico(self):
        assert build_pricebook_id("ml", "3333") is None

    def test_sem_runrun_id_nao_gera(self):
        assert build_pricebook_id("natura", "") is None
        assert build_pricebook_id("natura", None) is None


class TestParseRunrunId:
    def test_extrai_de_id_natura(self):
        assert parse_runrun_id("NAT-RR1111") == "1111"

    def test_extrai_de_id_avon(self):
        assert parse_runrun_id("AVN-RR2222") == "2222"

    def test_case_insensitive(self):
        assert parse_runrun_id("nat-rr1111") == "1111"

    def test_id_manual_nao_bate_no_padrao(self):
        assert parse_runrun_id("CB-PROMO-VERAO") is None

    def test_vazio(self):
        assert parse_runrun_id("") is None


# ─── compute_status — bordas exatas ─────────────────────────────────────

class TestComputeStatus:
    def test_sem_janela_e_ativa(self):
        assert compute_status(_record()) == "ativa"

    def test_dentro_da_janela_e_ativa(self):
        now = datetime(2026, 9, 10, tzinfo=timezone.utc)
        rec = _record(
            online_from="2026-09-01T00:00:00.000Z",
            online_to="2026-09-14T23:59:59.000Z",
        )
        assert compute_status(rec, now) == "ativa"

    def test_antes_da_janela_e_agendada(self):
        now = datetime(2026, 8, 31, tzinfo=timezone.utc)
        rec = _record(online_from="2026-09-01T00:00:00.000Z")
        assert compute_status(rec, now) == "agendada"

    def test_exatamente_no_online_from_e_ativa(self):
        now = datetime(2026, 9, 1, 0, 0, 0, tzinfo=timezone.utc)
        rec = _record(online_from="2026-09-01T00:00:00.000Z")
        assert compute_status(rec, now) == "ativa"

    def test_depois_da_janela_e_expirada(self):
        now = datetime(2026, 9, 15, tzinfo=timezone.utc)
        rec = _record(online_to="2026-09-14T23:59:59.000Z")
        assert compute_status(rec, now) == "expirada"

    def test_exatamente_no_online_to_e_ativa(self):
        now = datetime(2026, 9, 14, 23, 59, 59, tzinfo=timezone.utc)
        rec = _record(online_to="2026-09-14T23:59:59.000Z")
        assert compute_status(rec, now) == "ativa"

    def test_um_segundo_depois_do_online_to_e_expirada(self):
        now = datetime(2026, 9, 15, 0, 0, 0, tzinfo=timezone.utc)
        rec = _record(online_to="2026-09-14T23:59:59.000Z")
        assert compute_status(rec, now) == "expirada"

    def test_ended_at_vence_qualquer_janela(self):
        now = datetime(2026, 9, 10, tzinfo=timezone.utc)
        rec = _record(
            online_from="2026-09-01T00:00:00.000Z",
            online_to="2026-09-14T23:59:59.000Z",
            ended_at="2026-09-05T00:00:00.000Z",
        )
        assert compute_status(rec, now) == "encerrada"


# ─── resolve — regra de reconhecimento (Etapa 3, seção 8) ───────────────

class TestResolve:
    def test_sem_nenhum_registro_para_a_aba_e_none(self):
        result = resolve("LISTA_99", None, "natura", "natura", [])
        assert result.confidence == MatchConfidence.NONE
        assert result.record is None
        assert not result.is_risky

    def test_aba_marca_loja_e_lp_label_batendo_e_exact(self):
        rec = _record()
        result = resolve("lista_48", "Favoritos RR18/21", "natura", "natura", [rec])
        assert result.confidence == MatchConfidence.EXACT
        assert result.record is rec
        assert not result.is_risky

    def test_exact_sem_lp_label_informado_nao_e_risco(self):
        rec = _record(lp_label=None)
        result = resolve("LISTA_48", None, "natura", "natura", [rec])
        assert result.confidence == MatchConfidence.EXACT
        assert not result.is_risky

    def test_lp_label_divergente_vira_risco_mesmo_com_marca_e_loja_batendo(self):
        rec = _record(lp_label="Favoritos RR18/21")
        result = resolve("LISTA_48", "Outra Campanha", "natura", "natura", [rec])
        assert result.confidence == MatchConfidence.EXACT
        assert result.is_risky
        assert "rótulo LP divergente do registro" in result.reasons

    def test_registro_expirado_vira_risco(self):
        rec = _record(online_to="2020-01-01T00:00:00.000Z")
        result = resolve("LISTA_48", None, "natura", "natura", [rec])
        assert result.is_risky
        assert "registro expirado" in result.reasons

    def test_registro_com_mais_de_90_dias_vira_risco(self):
        old = (datetime.now(timezone.utc) - timedelta(days=91)).strftime("%Y-%m-%dT%H:%M:%S.000Z")
        rec = _record(updated_at=old)
        result = resolve("LISTA_48", None, "natura", "natura", [rec])
        assert result.is_risky
        assert "registro com mais de 90 dias" in result.reasons

    def test_registro_com_exatos_90_dias_nao_e_risco(self):
        edge = (datetime.now(timezone.utc) - timedelta(days=90)).strftime("%Y-%m-%dT%H:%M:%S.000Z")
        rec = _record(updated_at=edge)
        result = resolve("LISTA_48", None, "natura", "natura", [rec])
        assert not result.is_risky

    def test_mesma_aba_registrada_para_outra_marca_loja_e_sheet_only(self):
        # D4: a loja participa do casamento — sem isso, cairia em "ambíguo"
        # sem explicar que na verdade é um pricebook de outra loja/marca.
        rec = _record(brand="avon", loja="avon")
        result = resolve("LISTA_48", None, "natura", "natura", [rec])
        assert result.confidence == MatchConfidence.SHEET_ONLY
        assert result.is_risky

    def test_mesma_aba_natura_loja_e_natura_ml_nao_e_ambiguo(self):
        # A mesma aba pode virar pricebook na loja da marca E na Minha Loja
        # (CB) — são dois registros distintos, não uma ambiguidade (D4).
        rec_natura = _record(pricebook_id="NAT-RR1111", brand="natura", loja="natura")
        rec_ml = _record(pricebook_id="NAT-RR1111-CB", brand="natura", loja="ml")
        result = resolve("LISTA_48", None, "natura", "natura", [rec_natura, rec_ml])
        assert result.confidence == MatchConfidence.EXACT
        assert result.record is rec_natura

    def test_dois_registros_para_mesma_aba_marca_e_loja_e_ambiguous(self):
        rec_a = _record(pricebook_id="NAT-RR1111")
        rec_b = _record(pricebook_id="NAT-RR2222")
        result = resolve("LISTA_48", None, "natura", "natura", [rec_a, rec_b])
        assert result.confidence == MatchConfidence.AMBIGUOUS
        assert result.is_risky

    def test_dois_registros_para_mesma_aba_em_outras_marcas_e_ambiguous(self):
        rec_a = _record(pricebook_id="AVN-RR1111", brand="avon", loja="avon")
        rec_b = _record(pricebook_id="AVN-RR1111-CB", brand="avon", loja="ml")
        result = resolve("LISTA_48", None, "natura", "natura", [rec_a, rec_b])
        assert result.confidence == MatchConfidence.AMBIGUOUS
        assert result.is_risky

    def test_normalizacao_do_nome_da_aba_e_usada_no_casamento(self):
        rec = _record(sheet_name="  lista_48  ")
        result = resolve("LISTA_48", None, "natura", "natura", [rec])
        assert result.confidence == MatchConfidence.EXACT


# ─── InMemoryRegistryStore ───────────────────────────────────────────────

class TestInMemoryRegistryStore:
    def test_upsert_e_load_round_trip(self):
        store = InMemoryRegistryStore()
        rec = _record()
        store.upsert(rec)
        assert store.load() == [rec]

    def test_upsert_por_pricebook_id_nao_duplica(self):
        store = InMemoryRegistryStore()
        store.upsert(_record(sku_count=128))
        store.upsert(_record(sku_count=118))
        loaded = store.load()
        assert len(loaded) == 1
        assert loaded[0].sku_count == 118

    def test_mark_ended(self):
        store = InMemoryRegistryStore([_record()])
        store.mark_ended("NAT-RR1111", "2026-09-20T00:00:00.000Z")
        assert store.load()[0].ended_at == "2026-09-20T00:00:00.000Z"

    def test_mark_ended_de_id_inexistente_nao_quebra(self):
        store = InMemoryRegistryStore()
        store.mark_ended("NAT-RR9999", "2026-09-20T00:00:00.000Z")
        assert store.load() == []

    def test_delete(self):
        store = InMemoryRegistryStore([_record()])
        store.delete("NAT-RR1111")
        assert store.load() == []

    def test_delete_de_id_inexistente_nao_quebra(self):
        store = InMemoryRegistryStore()
        store.delete("NAT-RR9999")
        assert store.load() == []


# ─── JsonFileRegistryStore ───────────────────────────────────────────────

class TestJsonFileRegistryStore:
    def test_arquivo_ausente_devolve_lista_vazia(self, tmp_path: Path):
        store = JsonFileRegistryStore(tmp_path / "registry.json")
        assert store.load() == []

    def test_upsert_e_load_round_trip(self, tmp_path: Path):
        path = tmp_path / "registry.json"
        store = JsonFileRegistryStore(path)
        rec = _record()
        store.upsert(rec)

        reloaded = JsonFileRegistryStore(path).load()
        assert reloaded == [rec]

    def test_upsert_por_pricebook_id_atualiza_no_lugar(self, tmp_path: Path):
        path = tmp_path / "registry.json"
        store = JsonFileRegistryStore(path)
        store.upsert(_record(sku_count=128))
        store.upsert(_record(sku_count=118))
        loaded = store.load()
        assert len(loaded) == 1
        assert loaded[0].sku_count == 118

    def test_escrita_e_atomica_nao_deixa_tmp_para_tras(self, tmp_path: Path):
        path = tmp_path / "registry.json"
        JsonFileRegistryStore(path).upsert(_record())
        assert path.exists()
        assert not path.with_name(path.name + ".tmp").exists()

    def test_json_corrompido_vira_bak_e_recomeca_vazio(self, tmp_path: Path):
        path = tmp_path / "registry.json"
        path.write_text("{ isso não é json válido", encoding="utf-8")

        store = JsonFileRegistryStore(path)
        assert store.load() == []
        assert path.with_name(path.name + ".bak").exists()

    def test_schema_version_desconhecido_vira_bak_e_recomeca_vazio(self, tmp_path: Path):
        path = tmp_path / "registry.json"
        path.write_text(json.dumps({"schema_version": 999, "records": []}), encoding="utf-8")

        store = JsonFileRegistryStore(path)
        assert store.load() == []
        assert path.with_name(path.name + ".bak").exists()

    def test_mark_ended(self, tmp_path: Path):
        path = tmp_path / "registry.json"
        store = JsonFileRegistryStore(path)
        store.upsert(_record())
        store.mark_ended("NAT-RR1111", "2026-09-20T00:00:00.000Z")
        assert store.load()[0].ended_at == "2026-09-20T00:00:00.000Z"

    def test_cria_pasta_pai_se_necessario(self, tmp_path: Path):
        path = tmp_path / "sub" / "dir" / "registry.json"
        JsonFileRegistryStore(path).upsert(_record())
        assert path.exists()

    def test_delete(self, tmp_path: Path):
        path = tmp_path / "registry.json"
        store = JsonFileRegistryStore(path)
        store.upsert(_record())
        store.delete("NAT-RR1111")
        assert store.load() == []

    def test_delete_de_id_inexistente_nao_quebra(self, tmp_path: Path):
        path = tmp_path / "registry.json"
        store = JsonFileRegistryStore(path)
        store.upsert(_record())
        store.delete("NAT-RR9999")
        assert len(store.load()) == 1


# ─── merge_by_pricebook_id ────────────────────────────────────────────────

class TestMergeByPricebookId:
    def test_junta_fontes_sem_sobreposicao(self):
        local = [_record(pricebook_id="NAT-RR1111")]
        shared = [_record(pricebook_id="NAT-RR2222")]
        merged = merge_by_pricebook_id(local, shared)
        assert {r.pricebook_id for r in merged} == {"NAT-RR1111", "NAT-RR2222"}

    def test_mesmo_id_em_duas_fontes_fica_com_o_updated_at_mais_novo(self):
        old = _record(sku_count=128, updated_at="2026-08-01T00:00:00.000Z")
        new = _record(sku_count=118, updated_at="2026-09-01T00:00:00.000Z")
        merged = merge_by_pricebook_id([old], [new])
        assert len(merged) == 1
        assert merged[0].sku_count == 118

    def test_ordem_dos_argumentos_nao_importa_quando_ha_updated_at(self):
        old = _record(sku_count=128, updated_at="2026-08-01T00:00:00.000Z")
        new = _record(sku_count=118, updated_at="2026-09-01T00:00:00.000Z")
        merged = merge_by_pricebook_id([new], [old])
        assert merged[0].sku_count == 118

    def test_sem_updated_at_nao_quebra(self):
        a = _record(updated_at=None)
        merged = merge_by_pricebook_id([a])
        assert merged == [a]


# ─── SyncedFolderRegistryStore (P5, pasta do Google Drive) ────────────────

class TestSyncedFolderRegistryStore:
    def test_pasta_ausente_devolve_lista_vazia(self, tmp_path: Path):
        store = SyncedFolderRegistryStore(tmp_path / "nao_existe")
        assert store.load() == []

    def test_upsert_cria_um_arquivo_por_pricebook_id(self, tmp_path: Path):
        folder = tmp_path / "drive"
        store = SyncedFolderRegistryStore(folder)
        store.upsert(_record(pricebook_id="NAT-RR1111"))
        store.upsert(_record(pricebook_id="NAT-RR2222"))
        assert sorted(p.name for p in folder.glob("*.json")) == ["NAT-RR1111.json", "NAT-RR2222.json"]

    def test_upsert_por_pricebook_id_atualiza_no_lugar_sem_duplicar(self, tmp_path: Path):
        store = SyncedFolderRegistryStore(tmp_path / "drive")
        store.upsert(_record(sku_count=128))
        store.upsert(_record(sku_count=118))
        loaded = store.load()
        assert len(loaded) == 1
        assert loaded[0].sku_count == 118

    def test_duas_campanhas_diferentes_nao_disputam_o_mesmo_arquivo(self, tmp_path: Path):
        folder = tmp_path / "drive"
        store = SyncedFolderRegistryStore(folder)
        store.upsert(_record(pricebook_id="NAT-RR1111", sheet_name="LISTA_A"))
        store.upsert(_record(pricebook_id="NAT-RR2222", sheet_name="LISTA_B"))
        names = {r.pricebook_id: r.sheet_name for r in store.load()}
        assert names == {"NAT-RR1111": "LISTA_A", "NAT-RR2222": "LISTA_B"}

    def test_mark_ended(self, tmp_path: Path):
        store = SyncedFolderRegistryStore(tmp_path / "drive")
        store.upsert(_record())
        store.mark_ended("NAT-RR1111", "2026-09-20T00:00:00.000Z")
        assert store.load()[0].ended_at == "2026-09-20T00:00:00.000Z"

    def test_delete(self, tmp_path: Path):
        store = SyncedFolderRegistryStore(tmp_path / "drive")
        store.upsert(_record())
        store.delete("NAT-RR1111")
        assert store.load() == []

    def test_delete_de_id_inexistente_nao_quebra(self, tmp_path: Path):
        store = SyncedFolderRegistryStore(tmp_path / "drive")
        store.delete("NAT-RR9999")
        assert store.load() == []

    def test_arquivo_individual_corrompido_e_isolado_sem_derrubar_os_demais(self, tmp_path: Path):
        folder = tmp_path / "drive"
        store = SyncedFolderRegistryStore(folder)
        store.upsert(_record(pricebook_id="NAT-RR1111"))
        (folder / "NAT-RR2222.json").write_text("{ nao e json valido", encoding="utf-8")

        loaded = store.load()
        assert [r.pricebook_id for r in loaded] == ["NAT-RR1111"]
        assert (folder / "NAT-RR2222.json.bak").exists()

    def test_pricebook_id_com_caractere_incomum_vira_nome_de_arquivo_seguro(self, tmp_path: Path):
        folder = tmp_path / "drive"
        store = SyncedFolderRegistryStore(folder)
        store.upsert(_record(pricebook_id="CB/Promo Verão!"))
        files = list(folder.glob("*.json"))
        assert len(files) == 1
        assert store.load()[0].pricebook_id == "CB/Promo Verão!"


# ─── write_xlsx_snapshot ───────────────────────────────────────────────────

class TestWriteXlsxSnapshot:
    def test_gera_cabecalho_e_uma_linha_por_registro(self, tmp_path: Path):
        path = tmp_path / "drive" / "registry.xlsx"
        records = [_record(pricebook_id="NAT-RR1111"), _record(pricebook_id="NAT-RR2222")]
        write_xlsx_snapshot(records, path)

        wb = openpyxl.load_workbook(path)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0][0] == "pricebook_id"
        assert len(rows) == 3  # cabeçalho + 2 registros

    def test_lista_vazia_gera_so_o_cabecalho(self, tmp_path: Path):
        path = tmp_path / "registry.xlsx"
        write_xlsx_snapshot([], path)
        wb = openpyxl.load_workbook(path)
        rows = list(wb.active.iter_rows(values_only=True))
        assert len(rows) == 1

    def test_escrita_e_atomica_nao_deixa_tmp_para_tras(self, tmp_path: Path):
        path = tmp_path / "registry.xlsx"
        write_xlsx_snapshot([_record()], path)
        assert path.exists()
        assert not path.with_name(path.name + ".tmp").exists()

    def test_sobrescreve_snapshot_anterior(self, tmp_path: Path):
        path = tmp_path / "registry.xlsx"
        write_xlsx_snapshot([_record(pricebook_id="NAT-RR1111")], path)
        write_xlsx_snapshot([_record(pricebook_id="NAT-RR2222")], path)
        wb = openpyxl.load_workbook(path)
        rows = list(wb.active.iter_rows(values_only=True))
        assert len(rows) == 2  # cabeçalho + 1 registro (não acumula do anterior)
