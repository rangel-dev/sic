"""
Limpeza de Campanhas — motor (BRD-015).

Lê as 3 exportações do Business Manager (promoções/campanhas, cupons e
customer groups), classifica cada item em Protegido / Elegível / Revisar e,
só depois de confirmação explícita na tela, gera os XMLs de exclusão e de
rollback.

Regra central (decisão do negócio, BRD-015 seção 14): campanha vencida há
mais que a carência, ou sem data de fim e desativada, sai; promoção, cupom
ou customer group sem vínculo com uma campanha que fica também sai — já na
primeira rodada. O que protege contra erro é a cascata (item ligado a
campanha que fica nunca sai), o rollback completo (coupon-codes e
group-assignment, conferidos antes de liberar) e as travas de coerência e
de volume.

O SIC nunca importa nada no SFCC: só gera arquivos para o operador importar.
"""
from __future__ import annotations

import copy
import os
import re
import shutil
import tempfile
import xml.etree.ElementTree as ET
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Callable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ── Namespaces SFCC ──────────────────────────────────────────────────────────
PROMOTIONS_NS = "http://www.demandware.com/xml/impex/promotion/2008-01-31"
COUPONS_NS = "http://www.demandware.com/xml/impex/coupon/2008-06-17"
CUSTOMER_GROUPS_NS = "http://www.demandware.com/xml/impex/customergroup/2007-06-30"

KIND_PROMOTIONS = "promotions"
KIND_COUPONS = "coupons"
KIND_CUSTOMER_GROUPS = "customer_groups"

KIND_BY_NAMESPACE = {
    PROMOTIONS_NS: KIND_PROMOTIONS,
    COUPONS_NS: KIND_COUPONS,
    CUSTOMER_GROUPS_NS: KIND_CUSTOMER_GROUPS,
}
KIND_LABELS = {
    KIND_PROMOTIONS: "Promoções e campanhas",
    KIND_COUPONS: "Cupons",
    KIND_CUSTOMER_GROUPS: "Customer groups",
}

# ── Entidades e status ───────────────────────────────────────────────────────
ENT_CAMPAIGNS = "campaigns"
ENT_PROMOTIONS = "promotions"
ENT_COUPONS = "coupons"
ENT_GROUPS = "customer_groups"
ENTITIES = (ENT_CAMPAIGNS, ENT_PROMOTIONS, ENT_COUPONS, ENT_GROUPS)
ENTITY_LABELS = {
    ENT_CAMPAIGNS: "Campanhas",
    ENT_PROMOTIONS: "Promoções",
    ENT_COUPONS: "Cupons",
    ENT_GROUPS: "Customer groups",
}

PROTEGIDO = "Protegido"
ELEGIVEL = "Elegível"            # entra no XML de exclusão
REVISAR = "Revisar"              # nunca entra em XML; decisão humana
STATUSES = (PROTEGIDO, ELEGIVEL, REVISAR)

# Sugestões iniciais; o operador pode digitar qualquer loja LATAM (N6)
DEFAULT_LOJAS = ("Natura Brasil", "Avon Brasil", "Minha Loja (CB)")

# ── Parâmetros (BRD-015 — respostas do negócio às pendências N1 a N6) ───────
GRACE_DAYS = 30                # padrão da régua (N3) — só para campanhas vencidas
GRACE_MIN_DAYS, GRACE_MAX_DAYS = 7, 365
DEFAULT_PROTECTED_TERMS: tuple[str, ...] = ()   # N2: nenhum termo por padrão
SYSTEM_GROUPS = {"everyone", "registered", "unregistered"}   # grupos padrão do SFCC
INFRA_GROUP_TERMS = ("dont_remove", "dontremove")            # o próprio nome pede para não apagar
COHERENCE_MAX_MISSING = 0.10   # trava 7.1
VOLUME_BRAKE = 0.30            # trava 7.2


def loja_key(name: str) -> str:
    """Chave estável da loja para o nome da pasta de saída: sem acento,
    minúscula, só letras/números/hífen. "Natura Chile" → "natura-chile"."""
    import unicodedata
    plain = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", plain.lower()).strip("-")


ProgressFn = Callable[[int, str], None]


class LimpezaError(Exception):
    """Erro de entrada ou de conferência — nada é gerado."""


def _local(tag: str) -> str:
    return tag.split("}", 1)[1] if "}" in tag else tag


def parse_sfcc_date(text: Optional[str]) -> tuple[Optional[datetime], bool]:
    """Retorna (data, ilegível). Vazio → (None, False). Sem fuso → UTC.
    Ilegível → (None, True): o chamador manda o item para Revisar."""
    if not text or not text.strip():
        return None, False
    try:
        dt = datetime.fromisoformat(text.strip().replace("Z", "+00:00"))
    except ValueError:
        return None, True
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt, False


def _fmt_dt(dt: Optional[datetime]) -> str:
    return dt.strftime("%d/%m/%Y %H:%M") if dt else ""


# ─────────────────────────────────────────────────────────────────────────────
#  Reconhecimento dos arquivos pelo namespace (nunca pelo nome)
# ─────────────────────────────────────────────────────────────────────────────
def detect_kind(path: str) -> Optional[str]:
    """Tipo do export pelo namespace do elemento raiz, ou None."""
    try:
        for _ev, el in ET.iterparse(path, events=("start",)):
            ns = el.tag[1:].split("}", 1)[0] if el.tag.startswith("{") else ""
            return KIND_BY_NAMESPACE.get(ns)
    except (ET.ParseError, OSError):
        return None
    return None


@dataclass
class FileSelection:
    paths: dict[str, str] = field(default_factory=dict)   # kind -> path
    rejected: list[tuple[str, str]] = field(default_factory=list)  # (path, motivo)

    @property
    def complete(self) -> bool:
        return all(k in self.paths for k in KIND_BY_NAMESPACE.values())


def classify_files(paths: list[str], current: Optional[dict[str, str]] = None) -> FileSelection:
    """Encaixa cada arquivo no seu tipo. Arquivo desconhecido ou um segundo
    arquivo do mesmo tipo é recusado (trava 7.3)."""
    sel = FileSelection(paths=dict(current or {}))
    for p in paths:
        kind = detect_kind(p)
        if kind is None:
            sel.rejected.append((p, "não é um export de promoções, cupons ou customer groups do SFCC"))
        elif kind in sel.paths and os.path.abspath(sel.paths[kind]) != os.path.abspath(p):
            sel.rejected.append((p, f"já existe um arquivo de {KIND_LABELS[kind].lower()} selecionado"))
        else:
            sel.paths[kind] = p
    return sel


# ─────────────────────────────────────────────────────────────────────────────
#  Resultado
# ─────────────────────────────────────────────────────────────────────────────
@dataclass
class Item:
    entity: str
    id: str
    status: str = PROTEGIDO
    reason: str = ""
    info: dict = field(default_factory=dict)


@dataclass
class SimulationResult:
    loja: str            # nome exibido ("Natura Chile")
    loja_key: str        # usada no nome da pasta de saída ("natura-chile")
    now: datetime
    grace_days: int
    files: dict[str, str]
    items: dict[str, dict[str, Item]]
    coherence_ok: bool = True
    coherence_msg: str = ""
    missing_coupons: int = 0
    cited_coupons: int = 0
    missing_groups: int = 0
    cited_groups: int = 0
    volume_alerts: list[str] = field(default_factory=list)

    def counts(self, entity: str) -> Counter:
        return Counter(it.status for it in self.items[entity].values())

    def ids(self, entity: str, status: str) -> set[str]:
        return {i for i, it in self.items[entity].items() if it.status == status}

    @property
    def affected_total(self) -> int:
        """Itens que entram nos XMLs de exclusão."""
        return sum(self.counts(e)[ELEGIVEL] for e in ENTITIES)

    @property
    def can_export(self) -> bool:
        return self.coherence_ok and self.affected_total > 0


@dataclass
class ExportResult:
    folder: str
    files: list[str]


# ─────────────────────────────────────────────────────────────────────────────
#  Motor
# ─────────────────────────────────────────────────────────────────────────────
class LimpezaEngine:
    def __init__(
        self,
        files: dict[str, str],
        loja: str,
        protected_terms: Optional[list[str]] = None,
        grace_days: int = GRACE_DAYS,
        now: Optional[datetime] = None,
        progress_callback: Optional[ProgressFn] = None,
    ):
        missing = [KIND_LABELS[k] for k in KIND_BY_NAMESPACE.values() if k not in files]
        if missing:
            raise LimpezaError("Faltam os arquivos: " + ", ".join(missing) + ".")
        for kind, path in files.items():
            if detect_kind(path) != kind:
                raise LimpezaError(f"O arquivo {Path(path).name} não é um export de {KIND_LABELS[kind].lower()}.")
        self.loja = " ".join((loja or "").split())
        self.loja_key = loja_key(self.loja)
        if not self.loja_key:
            raise LimpezaError("Informe a loja da limpeza.")
        if not GRACE_MIN_DAYS <= grace_days <= GRACE_MAX_DAYS:
            raise LimpezaError(f"A carência precisa ficar entre {GRACE_MIN_DAYS} e {GRACE_MAX_DAYS} dias.")
        self.files = dict(files)
        terms = DEFAULT_PROTECTED_TERMS if protected_terms is None else protected_terms
        self.protected_terms = [t.strip().lower() for t in terms if t and t.strip()]
        self.grace = timedelta(days=grace_days)
        self.grace_days = grace_days
        self.now = now or datetime.now(timezone.utc)
        self._progress = progress_callback or (lambda _p, _m: None)

        self.campaigns: dict[str, dict] = {}
        self.promotions: dict[str, dict] = {}
        self.coupons: dict[str, dict] = {}
        self.groups: dict[str, dict] = {}
        self._promo_root: Optional[ET.Element] = None
        self._coupon_campaigns: dict[str, set[str]] = {}
        self._group_campaigns: dict[str, set[str]] = {}

    def set_progress_callback(self, fn: Optional[ProgressFn]) -> None:
        self._progress = fn or (lambda _p, _m: None)

    # ── Leitura ───────────────────────────────────────────────────────────
    def _is_protected_term(self, *texts: str) -> Optional[str]:
        low = " ".join(t.lower() for t in texts if t)
        return next((t for t in self.protected_terms if t in low), None)

    def _load_promotions(self):
        self._progress(5, "Lendo promoções e campanhas…")
        root = ET.parse(self.files[KIND_PROMOTIONS]).getroot()
        self._promo_root = root
        coupon_links: dict[str, set[str]] = {}
        group_links: dict[str, set[str]] = {}

        def link(target: dict[str, set[str]], key: str, camp: str):
            if key:
                target.setdefault(key, set()).add(camp)

        for el in root:
            name = _local(el.tag)
            if name == "campaign":
                cid = el.attrib.get("campaign-id", "").strip()
                if not cid:
                    continue
                start, bad_start = parse_sfcc_date(el.findtext("{*}start-date"))
                end, bad_end = parse_sfcc_date(el.findtext("{*}end-date"))
                self.campaigns[cid] = {
                    "description": (el.findtext("{*}description") or "").strip(),
                    "enabled": (el.findtext("{*}enabled-flag") or "true").strip().lower() == "true",
                    "start": start, "end": end, "invalid_date": bad_start or bad_end,
                    "start_txt": el.findtext("{*}start-date") or "",
                    "end_txt": el.findtext("{*}end-date") or "",
                    "promotions": set(),
                }
                for c in el.iterfind(".//{*}coupons/{*}coupon"):
                    link(coupon_links, c.attrib.get("coupon-id", "").strip(), cid)
                for g in el.iterfind(".//{*}customer-groups/{*}customer-group"):
                    link(group_links, g.attrib.get("group-id", "").strip(), cid)
            elif name == "promotion":
                pid = el.attrib.get("promotion-id", "").strip()
                if pid:
                    self.promotions[pid] = {
                        "name": (el.findtext("{*}name") or "").strip(),
                        "enabled": (el.findtext("{*}enabled-flag") or "true").strip().lower() == "true",
                        "campaigns": set(),
                    }
        for el in root:
            if _local(el.tag) != "promotion-campaign-assignment":
                continue
            cid = el.attrib.get("campaign-id", "").strip()
            pid = el.attrib.get("promotion-id", "").strip()
            if pid in self.promotions and cid:
                self.promotions[pid]["campaigns"].add(cid)
            if cid in self.campaigns and pid:
                self.campaigns[cid]["promotions"].add(pid)
            for c in el.iterfind(".//{*}coupons/{*}coupon"):
                link(coupon_links, c.attrib.get("coupon-id", "").strip(), cid)
            for g in el.iterfind(".//{*}customer-groups/{*}customer-group"):
                link(group_links, g.attrib.get("group-id", "").strip(), cid)
        self._coupon_campaigns = coupon_links
        self._group_campaigns = group_links

    def _load_coupons(self):
        """Streaming: os códigos (milhões) são só contados e descartados na
        hora — o pacote original os mantinha em memória (BRD-015, 5.4)."""
        self._progress(30, "Lendo cupons (pode levar alguns segundos)…")
        codes: Counter = Counter()
        root = block = None
        block_id = ""
        depth = 0
        for ev, el in ET.iterparse(self.files[KIND_COUPONS], events=("start", "end")):
            if ev == "start":
                depth += 1
                if root is None:
                    root = el
                elif depth == 2 and _local(el.tag) == "coupon-codes":
                    block, block_id = el, el.attrib.get("coupon-id", "").strip()
                continue
            depth -= 1
            name = _local(el.tag)
            if depth == 2 and block is not None and name == "code":
                codes[block_id] += 1
                del block[-1]
            elif depth == 1:
                if name == "coupon":
                    cid = el.attrib.get("coupon-id", "").strip()
                    if cid:
                        kinds = [k for k in ("single-code", "multiple-codes", "system-codes")
                                 if el.find("{*}" + k) is not None]
                        self.coupons[cid] = {
                            "description": (el.findtext("{*}description") or "").strip(),
                            "enabled": (el.findtext("{*}enabled-flag") or "true").strip().lower() == "true",
                            "code_type": kinds[0] if kinds else "",
                        }
                block = None
                root.clear()
        for cid, info in self.coupons.items():
            info["codes"] = codes.get(cid, 0)
            info["campaigns"] = self._coupon_campaigns.get(cid, set())

    def _load_groups(self):
        self._progress(50, "Lendo customer groups…")
        members: Counter = Counter()
        root = None
        depth = 0
        for ev, el in ET.iterparse(self.files[KIND_CUSTOMER_GROUPS], events=("start", "end")):
            if ev == "start":
                depth += 1
                root = root if root is not None else el
                continue
            depth -= 1
            if depth != 1:
                continue
            name = _local(el.tag)
            if name == "customer-group":
                gid = el.attrib.get("group-id", "").strip()
                if gid:
                    self.groups[gid] = {
                        "description": (el.findtext("{*}description") or "").strip(),
                        "dynamic": el.find("{*}membership-rule") is not None,
                    }
            elif name == "group-assignment":
                members[el.attrib.get("group-id", "").strip()] += 1
            root.clear()
        for gid, info in self.groups.items():
            info["members"] = members.get(gid, 0)
            info["campaigns"] = self._group_campaigns.get(gid, set())

    # ── Simulação ─────────────────────────────────────────────────────────
    def simulate(self) -> SimulationResult:
        """Classifica tudo. Não grava nada: o resultado depende só dos 3
        arquivos, da data e da carência — igual em qualquer computador."""
        self._load_promotions()
        self._load_coupons()
        self._load_groups()
        self._progress(65, "Conferindo coerência entre os arquivos…")

        res = SimulationResult(
            loja=self.loja, loja_key=self.loja_key, now=self.now, grace_days=self.grace_days,
            files=dict(self.files), items={e: {} for e in ENTITIES},
        )
        self._check_coherence(res)

        self._progress(75, f"Classificando (carência de {self.grace_days} dias para campanhas vencidas)…")
        blocking = self._classify_campaigns(res)
        self._classify_promotions(res, blocking)
        self._classify_coupons(res, blocking)
        self._classify_groups(res, blocking)
        self._check_volume(res)
        self._progress(100, "Simulação concluída.")
        return res

    def _check_coherence(self, res: SimulationResult):
        """Trava 7.1: arquivos de lojas/datas diferentes fazem os itens
        citados pelas campanhas sumirem dos outros arquivos."""
        cited_c = {c for c in self._coupon_campaigns if c}
        cited_g = {g for g in self._group_campaigns if g}
        res.cited_coupons, res.missing_coupons = len(cited_c), len(cited_c - self.coupons.keys())
        res.cited_groups, res.missing_groups = len(cited_g), len(cited_g - self.groups.keys())
        problems = []
        if not self.campaigns:
            problems.append("o arquivo de promoções não tem nenhuma campanha")
        for label, missing, cited in (("cupons", res.missing_coupons, res.cited_coupons),
                                      ("customer groups", res.missing_groups, res.cited_groups)):
            if cited and missing > 1 and missing / cited > COHERENCE_MAX_MISSING:
                problems.append(
                    f"{missing} de {cited} {label} citados pelas campanhas não existem no arquivo de {label} "
                    f"({missing / cited:.0%}, limite {COHERENCE_MAX_MISSING:.0%})")
        if problems:
            res.coherence_ok = False
            res.coherence_msg = ("Os 3 arquivos não parecem ser da mesma loja e do mesmo momento: "
                                 + "; ".join(problems) + ". Exporte os 3 de novo, juntos, da mesma loja.")

    def _days(self, since: datetime) -> int:
        return max(0, (self.now - since).days)

    def _classify_campaigns(self, res) -> set[str]:
        """Retorna as campanhas que seguram a cascata (tudo que não é Elegível)."""
        for cid, c in self.campaigns.items():
            it = Item(ENT_CAMPAIGNS, cid, info={
                "Habilitada": "Sim" if c["enabled"] else "Não",
                "Início": _fmt_dt(c["start"]) or c["start_txt"],
                "Fim": _fmt_dt(c["end"]) or c["end_txt"] or "Sem data de fim",
                "Promoções": len(c["promotions"]),
                "Descrição": c["description"],
            })
            term = self._is_protected_term(cid)
            end = c["end"]
            if term:
                it.reason = f'Lista protegida ("{term}")'
            elif c["invalid_date"]:
                it.status = REVISAR
                it.reason = "Data de início ou de fim ilegível — o SIC não decide sozinho"
            elif end is not None:
                if end >= self.now:
                    it.reason = "Futura" if c["start"] and c["start"] > self.now else "Vigente"
                elif self.now - end < self.grace:
                    it.reason = f"Vencida há {self._days(end)} dias (carência de {self.grace_days})"
                else:
                    it.status, it.reason = ELEGIVEL, f"Vencida há {self._days(end)} dias"
            elif c["enabled"]:
                it.reason = "Sem data de fim e habilitada (perene)"
            else:
                it.status, it.reason = ELEGIVEL, "Sem data de fim e desativada"
            res.items[ENT_CAMPAIGNS][cid] = it
        return {i for i, it in res.items[ENT_CAMPAIGNS].items() if it.status != ELEGIVEL}

    def _linked(self, camps: set[str], blocking: set[str]) -> tuple[list[str], list[str]]:
        held = sorted(camps & blocking)
        unknown = sorted(camps - self.campaigns.keys())
        return held, unknown

    def _classify_promotions(self, res, blocking):
        for pid, p in self.promotions.items():
            camps = p["campaigns"]
            it = Item(ENT_PROMOTIONS, pid, info={
                "Nome": p["name"], "Habilitada": "Sim" if p["enabled"] else "Não",
                "Campanhas": ", ".join(sorted(camps)) or "Nenhuma",
            })
            held, unknown = self._linked(camps, blocking)
            term = self._is_protected_term(pid)
            if term:
                it.reason = f'Lista protegida ("{term}")'
            elif not camps:
                it.status, it.reason = ELEGIVEL, "Sem campanha"
            elif held:
                it.reason = "Ligada a campanha que fica: " + ", ".join(held)
            elif unknown:
                it.status, it.reason = REVISAR, "Ligada a campanha que não está no arquivo: " + ", ".join(unknown)
            else:
                it.status, it.reason = ELEGIVEL, "Todas as campanhas dela serão excluídas"
            res.items[ENT_PROMOTIONS][pid] = it

    def _classify_coupons(self, res, blocking):
        for cid, c in self.coupons.items():
            camps = c["campaigns"]
            it = Item(ENT_COUPONS, cid, info={
                "Habilitado": "Sim" if c["enabled"] else "Não", "Tipo de código": c["code_type"],
                "Códigos no export": c["codes"], "Campanhas": ", ".join(sorted(camps)) or "Nenhuma",
                "Descrição": c["description"],
            })
            held, unknown = self._linked(camps, blocking)
            term = self._is_protected_term(cid)
            res.items[ENT_COUPONS][cid] = it
            if term:
                it.reason = f'Lista protegida ("{term}")'
                continue
            if held:
                it.reason = "Ligado a campanha que fica: " + ", ".join(held)
                if not c["enabled"]:
                    it.reason += " — está desativado; reative no Business Manager se a campanha precisar dele"
                continue
            if unknown:
                it.status, it.reason = REVISAR, "Ligado a campanha que não está no arquivo: " + ", ".join(unknown)
                continue
            # Regra do negócio: cupom sem vínculo com campanha que fica é excluído
            it.status = ELEGIVEL
            it.reason = "Sem campanha" if not camps else "Todas as campanhas dele serão excluídas"

    def _classify_groups(self, res, blocking):
        for gid, g in self.groups.items():
            camps = g["campaigns"]
            it = Item(ENT_GROUPS, gid, info={
                "Tipo": "Dinâmico (regra)" if g["dynamic"] else "Estático",
                "Membros estáticos": g["members"], "Campanhas": ", ".join(sorted(camps)) or "Nenhuma",
                "Descrição": g["description"],
            })
            held, unknown = self._linked(camps, blocking)
            term = self._is_protected_term(gid)
            infra = next((t for t in INFRA_GROUP_TERMS if t in f"{gid} {g['description']}".lower()), None)
            # Decisão do negócio (N1/N2): fora os grupos padrão do SFCC e os que
            # pedem no nome para não serem removidos, todo grupo segue as
            # regras — o padrão do ID não importa.
            if gid.lower() in SYSTEM_GROUPS:
                it.reason = "Grupo padrão do SFCC"
            elif infra:
                it.reason = f'O nome pede para não remover ("{infra}")'
            elif term:
                it.reason = f'Lista protegida ("{term}")'
            elif held:
                it.reason = "Ligado a campanha que fica: " + ", ".join(held)
            elif unknown:
                it.status, it.reason = REVISAR, "Ligado a campanha que não está no arquivo: " + ", ".join(unknown)
            elif not camps:
                it.status, it.reason = ELEGIVEL, "Sem campanha"
            else:
                it.status, it.reason = ELEGIVEL, "Todas as campanhas dele serão excluídas"
            res.items[ENT_GROUPS][gid] = it

    def _check_volume(self, res: SimulationResult):
        """Trava 7.2: muito volume numa rodada exige confirmação reforçada."""
        for e in ENTITIES:
            total = len(res.items[e])
            n = res.counts(e)[ELEGIVEL]
            if total and n / total > VOLUME_BRAKE:
                res.volume_alerts.append(
                    f"{ENTITY_LABELS[e]}: {n} de {total} ({n / total:.0%}, limite {VOLUME_BRAKE:.0%})")

    # ── Exportação (só após confirmação na tela) ─────────────────────────
    def export(self, res: SimulationResult, output_dir: str) -> ExportResult:
        if not res.coherence_ok:
            raise LimpezaError(res.coherence_msg)
        if self._promo_root is None:
            raise LimpezaError("Rode a simulação antes de gerar os arquivos.")
        stamp = res.now.astimezone().strftime("%Y%m%d_%H%M")
        final = Path(output_dir) / f"LIMPEZA_{res.loja_key.upper()}_{stamp}"
        if final.exists():
            raise LimpezaError(f"A pasta {final.name} já existe. Escolha outra pasta ou aguarde um minuto.")
        tmp = Path(tempfile.mkdtemp(prefix=".limpeza_", dir=output_dir))
        try:
            self._progress(10, "Gerando XMLs de promoções e campanhas…")
            names = self._write_promotions(res, tmp)
            self._progress(35, "Gerando XMLs de cupons (com os códigos no rollback)…")
            names += self._write_coupons(res, tmp)
            self._progress(65, "Gerando XMLs de customer groups…")
            names += self._write_groups(res, tmp)
            self._progress(80, "Gerando a planilha da simulação…")
            build_workbook(res).save(tmp / "SIMULACAO_LIMPEZA.xlsx")
            names.append("SIMULACAO_LIMPEZA.xlsx")
            self._progress(90, "Conferindo os rollbacks…")
            verify_export(res, self, tmp)
            os.replace(tmp, final)
        except Exception:
            shutil.rmtree(tmp, ignore_errors=True)
            raise
        self._progress(100, "Arquivos gerados.")
        return ExportResult(str(final), [str(final / n) for n in names])

    @staticmethod
    def _open_xml(path: Path, root_tag: str, ns: str):
        f = open(path, "w", encoding="utf-8", newline="\n")
        f.write(f'<?xml version="1.0" encoding="UTF-8"?>\n<{root_tag} xmlns="{ns}">\n')
        return f

    @staticmethod
    def _ser(el: ET.Element, ns: str) -> str:
        """Serializa sem o namespace nas tags: o elemento raiz do arquivo já
        declara o namespace padrão, então elas o herdam. Atributos xml:lang
        continuam qualificados. Altera `el` — passe uma cópia se precisar."""
        prefix = "{" + ns + "}"
        for e in el.iter():
            if isinstance(e.tag, str) and e.tag.startswith(prefix):
                e.tag = e.tag[len(prefix):]
        el.tail = None
        return "    " + ET.tostring(el, encoding="unicode") + "\n"

    def _write_promotions(self, res, folder: Path) -> list[str]:
        camps, promos = res.ids(ENT_CAMPAIGNS, ELEGIVEL), res.ids(ENT_PROMOTIONS, ELEGIVEL)
        with self._open_xml(folder / FILE_PROMO_DELETE, "promotions", PROMOTIONS_NS) as f:
            for cid in sorted(camps):
                f.write(f'    <campaign campaign-id="{_attr(cid)}" mode="delete"/>\n')
            for pid in sorted(promos):
                f.write(f'    <promotion promotion-id="{_attr(pid)}" mode="delete"/>\n')
            f.write("</promotions>\n")
        # Rollback: campanhas → promoções → vínculos (ordem do schema)
        def wanted(el: ET.Element, kind: str) -> bool:
            camp = el.attrib.get("campaign-id", "").strip()
            promo = el.attrib.get("promotion-id", "").strip()
            if kind == "campaign":
                return camp in camps
            if kind == "promotion":
                return promo in promos
            return camp in camps or promo in promos

        with self._open_xml(folder / FILE_PROMO_ROLLBACK, "promotions", PROMOTIONS_NS) as f:
            for kind in ("campaign", "promotion", "promotion-campaign-assignment"):
                for el in self._promo_root:
                    if _local(el.tag) == kind and wanted(el, kind):
                        f.write(self._ser(_copy(el), PROMOTIONS_NS))
            f.write("</promotions>\n")
        return [FILE_PROMO_DELETE, FILE_PROMO_ROLLBACK]

    def _write_coupons(self, res, folder: Path) -> list[str]:
        to_delete = res.ids(ENT_COUPONS, ELEGIVEL)
        with self._open_xml(folder / FILE_COUPON_DELETE, "coupons", COUPONS_NS) as f:
            for cid in sorted(to_delete):
                f.write(f'    <coupon coupon-id="{_attr(cid)}" mode="delete"/>\n')
            f.write("</coupons>\n")
        rollback = self._open_xml(folder / FILE_COUPON_ROLLBACK, "coupons", COUPONS_NS)
        codes_tmp = folder / ".codes.tmp"
        codes_out = open(codes_tmp, "w", encoding="utf-8", newline="\n")
        try:
            root = block = None
            keep_block = False
            depth = 0
            for ev, el in ET.iterparse(self.files[KIND_COUPONS], events=("start", "end")):
                if ev == "start":
                    depth += 1
                    if root is None:
                        root = el
                    elif depth == 2 and _local(el.tag) == "coupon-codes":
                        block = el
                        keep_block = el.attrib.get("coupon-id", "").strip() in to_delete
                    continue
                depth -= 1
                name = _local(el.tag)
                if depth == 2 and block is not None and name == "code" and not keep_block:
                    del block[-1]
                elif depth == 1:
                    if name == "coupon":
                        if el.attrib.get("coupon-id", "").strip() in to_delete:
                            rollback.write(self._ser(_copy(el), COUPONS_NS))
                    elif name == "coupon-codes" and keep_block:
                        codes_out.write(self._ser(el, COUPONS_NS))
                    block, keep_block = None, False
                    root.clear()
            codes_out.close()
            # Schema: todas as definições de cupom antes dos blocos de códigos
            with open(codes_tmp, encoding="utf-8") as src:
                shutil.copyfileobj(src, rollback)
            rollback.write("</coupons>\n")
        finally:
            codes_out.close()
            rollback.close()
            codes_tmp.unlink(missing_ok=True)
        return [FILE_COUPON_DELETE, FILE_COUPON_ROLLBACK]

    def _write_groups(self, res, folder: Path) -> list[str]:
        to_delete = res.ids(ENT_GROUPS, ELEGIVEL)
        with self._open_xml(folder / FILE_GROUP_DELETE, "customer-groups", CUSTOMER_GROUPS_NS) as f:
            for gid in sorted(to_delete):
                f.write(f'    <customer-group group-id="{_attr(gid)}" mode="delete"/>\n')
            f.write("</customer-groups>\n")
        groups_xml: list[str] = []
        members_xml: list[str] = []
        root = None
        depth = 0
        for ev, el in ET.iterparse(self.files[KIND_CUSTOMER_GROUPS], events=("start", "end")):
            if ev == "start":
                depth += 1
                root = root if root is not None else el
                continue
            depth -= 1
            if depth != 1:
                continue
            name = _local(el.tag)
            if el.attrib.get("group-id", "").strip() in to_delete:
                if name == "customer-group":
                    groups_xml.append(self._ser(_copy(el), CUSTOMER_GROUPS_NS))
                elif name == "group-assignment":
                    members_xml.append(self._ser(_copy(el), CUSTOMER_GROUPS_NS))
            root.clear()
        # Sempre gerado, mesmo vazio (BRD-015, 5.4)
        with self._open_xml(folder / FILE_GROUP_ROLLBACK, "customer-groups", CUSTOMER_GROUPS_NS) as f:
            f.writelines(groups_xml)
            f.writelines(members_xml)
            f.write("</customer-groups>\n")
        return [FILE_GROUP_DELETE, FILE_GROUP_ROLLBACK]


# Nomes na ordem de importação (exclusão 1→3; rollback 1→3 é a ordem inversa)
FILE_PROMO_DELETE = "1_EXCLUIR_CAMPANHAS_PROMOCOES.xml"
FILE_COUPON_DELETE = "2_EXCLUIR_CUPONS.xml"
FILE_GROUP_DELETE = "3_EXCLUIR_CUSTOMER_GROUPS.xml"
FILE_GROUP_ROLLBACK = "ROLLBACK_1_CUSTOMER_GROUPS.xml"
FILE_COUPON_ROLLBACK = "ROLLBACK_2_CUPONS.xml"
FILE_PROMO_ROLLBACK = "ROLLBACK_3_CAMPANHAS_PROMOCOES.xml"

IMPORT_ORDER = (
    ("Limpeza", [FILE_PROMO_DELETE, FILE_COUPON_DELETE, FILE_GROUP_DELETE]),
    ("Rollback (só se precisar voltar atrás)", [FILE_GROUP_ROLLBACK, FILE_COUPON_ROLLBACK, FILE_PROMO_ROLLBACK]),
)


def _attr(value: str) -> str:
    return (value.replace("&", "&amp;").replace('"', "&quot;")
            .replace("<", "&lt;").replace(">", "&gt;"))


def _copy(el: ET.Element) -> ET.Element:
    return copy.deepcopy(el)


# ─────────────────────────────────────────────────────────────────────────────
#  Conferência dos arquivos gerados (BRD-015, 5.3 — CA-13)
# ─────────────────────────────────────────────────────────────────────────────
def _scan(path: Path, id_attr_by_tag: dict[str, str]) -> dict[str, Counter]:
    """Conta, por tag de 1º nível, quantas vezes cada id aparece; e os
    filhos <code> por bloco coupon-codes."""
    found: dict[str, Counter] = {t: Counter() for t in id_attr_by_tag}
    found["code"] = Counter()
    depth = 0
    root = None
    block_id = ""
    for ev, el in ET.iterparse(path, events=("start", "end")):
        if ev == "start":
            depth += 1
            root = root if root is not None else el
            if depth == 2 and _local(el.tag) == "coupon-codes":
                block_id = el.attrib.get("coupon-id", "")
            continue
        depth -= 1
        name = _local(el.tag)
        if depth == 2 and name == "code":
            found["code"][block_id] += 1
        elif depth == 1:
            if name in id_attr_by_tag:
                found[name][el.attrib.get(id_attr_by_tag[name], "")] += 1
            root.clear()
    return found


def verify_export(res: SimulationResult, engine: LimpezaEngine, folder: Path) -> None:
    problems: list[str] = []

    def same(label: str, got: Counter, want: set[str]):
        if set(got) != want or any(v != 1 for v in got.values()):
            extra, missing = set(got) - want, want - set(got)
            dup = [k for k, v in got.items() if v != 1]
            problems.append(f"{label}: {len(missing)} faltando, {len(extra)} a mais, {len(dup)} repetidos")

    camps, promos = res.ids(ENT_CAMPAIGNS, ELEGIVEL), res.ids(ENT_PROMOTIONS, ELEGIVEL)
    c_del = res.ids(ENT_COUPONS, ELEGIVEL)
    g_del = res.ids(ENT_GROUPS, ELEGIVEL)
    p_tags = {"campaign": "campaign-id", "promotion": "promotion-id"}

    d = _scan(folder / FILE_PROMO_DELETE, p_tags)
    same("Exclusão de campanhas", d["campaign"], camps)
    same("Exclusão de promoções", d["promotion"], promos)
    r = _scan(folder / FILE_PROMO_ROLLBACK, p_tags)
    same("Rollback de campanhas", r["campaign"], camps)
    same("Rollback de promoções", r["promotion"], promos)

    d = _scan(folder / FILE_COUPON_DELETE, {"coupon": "coupon-id"})
    same("Exclusão de cupons", d["coupon"], c_del)
    r = _scan(folder / FILE_COUPON_ROLLBACK, {"coupon": "coupon-id", "coupon-codes": "coupon-id"})
    same("Rollback de cupons", r["coupon"], c_del)
    for cid in c_del:
        want = engine.coupons[cid]["codes"]
        if r["code"].get(cid, 0) != want:
            problems.append(f"Rollback de cupons: códigos de um cupom não batem ({r['code'].get(cid, 0)} de {want})")
            break

    d = _scan(folder / FILE_GROUP_DELETE, {"customer-group": "group-id"})
    same("Exclusão de customer groups", d["customer-group"], g_del)
    r = _scan(folder / FILE_GROUP_ROLLBACK, {"customer-group": "group-id", "group-assignment": "group-id"})
    same("Rollback de customer groups", r["customer-group"], g_del)
    for gid in g_del:
        if r["group-assignment"].get(gid, 0) != engine.groups[gid]["members"]:
            problems.append("Rollback de customer groups: membros estáticos não batem")
            break

    if problems:
        raise LimpezaError("A conferência dos arquivos falhou e nada foi gerado: " + "; ".join(problems) + ".")


# ─────────────────────────────────────────────────────────────────────────────
#  Planilha da simulação
# ─────────────────────────────────────────────────────────────────────────────
_HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_STATUS_FILL = {PROTEGIDO: "D9EAD3", ELEGIVEL: "F4CCCC", REVISAR: "D0E0E3"}


def build_workbook(res: SimulationResult) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "LIMPEZA DE CAMPANHAS — SIMULAÇÃO"
    ws["A1"].font = Font(bold=True, size=14, color="1F497D")
    ws["A2"] = (f"Loja: {res.loja}  |  Execução: {res.now.astimezone().strftime('%d/%m/%Y %H:%M')}"
                f"  |  Carência: {res.grace_days} dias")
    ws.append([])
    ws.append(["Entidade", "Total", *STATUSES])
    for cell in ws[4]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
    for e in ENTITIES:
        c = res.counts(e)
        ws.append([ENTITY_LABELS[e], len(res.items[e]), *(c[s] for s in STATUSES)])
    ws.append([])
    ws.append(["Coerência entre os arquivos", "OK" if res.coherence_ok else "BLOQUEADO"])
    if res.coherence_msg:
        ws.append([res.coherence_msg])
    ws.append([f"Cupons citados pelas campanhas que não existem no arquivo: "
               f"{res.missing_coupons} de {res.cited_coupons}"])
    for a in res.volume_alerts:
        ws.append([f"Freio de volume: {a}"])
    ws.append([])
    ws.append(["Ordem de importação no Business Manager"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for title, names in IMPORT_ORDER:
        ws.append([title])
        for i, n in enumerate(names, 1):
            ws.append([f"   {i}. {n}"])
    ws.append([])
    ws.append(["O rollback não restaura o histórico de resgates dos cupons."])
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 18

    for e in ENTITIES:
        sheet = wb.create_sheet(ENTITY_LABELS[e])
        items = sorted(res.items[e].values(), key=lambda it: (STATUSES.index(it.status), it.id.lower()))
        extra = list(items[0].info) if items else []
        sheet.append(["ID", "Status", "Motivo", *extra])
        for cell in sheet[1]:
            cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
            cell.alignment = Alignment(vertical="center")
        for it in items:
            sheet.append([it.id, it.status, it.reason, *(it.info.get(k, "") for k in extra)])
            color = _STATUS_FILL[it.status]
            sheet.cell(sheet.max_row, 2).fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        for idx in range(1, sheet.max_column + 1):
            width = max((len(str(c.value or "")) for c in sheet[get_column_letter(idx)][:200]), default=10)
            sheet.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 12), 70)
        sheet.freeze_panes = "A2"
    return wb
